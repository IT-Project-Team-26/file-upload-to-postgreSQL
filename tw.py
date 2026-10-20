import os
import json
import time
import logging
from openpyxl import load_workbook
from twilio.rest import Client
from twilio.base.exceptions import TwilioRestException

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

client = Client(os.environ["TWILIO_ACCOUNT_SID"], os.environ["TWILIO_AUTH_TOKEN"])

FROM_NUMBER = "whatsapp:+14155238886"          # your sender
CONTENT_SID = "HXxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx"  # your APPROVED template SID
EXCEL_FILE  = "contacts.xlsx"
DELAY_SEC   = 1.0                              # pause between sends

wb = load_workbook(EXCEL_FILE)
ws = wb.active

headers = [c.value for c in ws[1]]
if "phone" not in headers:
    raise SystemExit("Excel must have a 'phone' column")

phone_idx = headers.index("phone")
var_cols = {h: i for i, h in enumerate(headers) if h.startswith("var")}

sent, failed = 0, 0

for row in ws.iter_rows(min_row=2, values_only=True):
    phone = row[phone_idx]
    if not phone:
        continue

    phone = str(phone).strip()
    if not phone.startswith("+"):
        log.warning("Skipping %s — missing country code", phone)
        failed += 1
        continue

    # Build variables: {"1": ..., "2": ...}
    variables = {}
    for i, (col, idx) in enumerate(sorted(var_cols.items()), start=1):
        variables[str(i)] = str(row[idx]) if row[idx] is not None else ""

    try:
        msg = client.messages.create(
            from_=FROM_NUMBER,
            to=f"whatsapp:{phone}",
            content_sid=CONTENT_SID,
            content_variables=json.dumps(variables)
        )
        log.info("Sent -> %s | SID=%s | status=%s", phone, msg.sid, msg.status)
        sent += 1
    except TwilioRestException as e:
        log.error("Failed -> %s | code=%s | msg=%s", phone, e.code, e.msg)
        failed += 1

    time.sleep(DELAY_SEC)

log.info("Done. Sent=%d Failed=%d", sent, failed)